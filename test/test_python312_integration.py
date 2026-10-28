"""Exercise installed libraries on Python 3.12 without provider API calls."""

import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch


class Python312IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        environment = patch.dict(os.environ, {
            "GROQ_API_KEY": "test-groq-key",
            "GROQ_MODEL": "openai/gpt-oss-20b",
            "GROQ_MAX_TOKENS": "1024",
            "GEMINI_API_KEY": "test-gemini-key",
            "GRADIO_ANALYTICS_ENABLED": "False",
            "GRADIO_SHARE": "false",
            "HF_HUB_OFFLINE": "1",
            "HF_HUB_DISABLE_TELEMETRY": "1",
            "ANONYMIZED_TELEMETRY": "False",
        })
        environment.start()
        cls.addClassCleanup(environment.stop)
        from config.settings import settings
        cls.settings = settings

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        for name, value in {
            "CACHE_DIR": str(self.root / "cache"),
            "CHROMA_DB_PATH": str(self.root / "chroma"),
        }.items():
            override = patch.object(self.settings, name, value)
            override.start()
            self.addCleanup(override.stop)

    def test_settings_load_dotenv_and_environment_precedence(self):
        from config.settings import Settings

        dotenv = self.root / ".env"
        dotenv.write_text(
            "GROQ_API_KEY=dotenv-groq\n",
            encoding="utf-8",
        )
        with patch.dict(os.environ, {}, clear=True):
            values = Settings(_env_file=dotenv)
            self.assertEqual(values.GROQ_API_KEY, "dotenv-groq")
            self.assertIsNone(values.GEMINI_API_KEY)
            os.environ["GROQ_API_KEY"] = "environment-groq"
            self.assertEqual(Settings(_env_file=dotenv).GROQ_API_KEY, "environment-groq")

    def test_gradio_filepaths_conversion_deduplication_and_cache(self):
        import app
        from document_processor.file_handler import DocumentProcessor

        path = self.root / "report.md"
        path.write_text("# Cases\n\nThe incident report records 42 cases.\n", encoding="utf-8")
        processor = DocumentProcessor()
        chunks = processor.process([str(path), path])
        self.assertTrue(chunks)
        self.assertIn("42 cases", " ".join(chunk.page_content for chunk in chunks))
        self.assertEqual(len(chunks), len({chunk.page_content for chunk in chunks}))
        with patch.object(processor, "_process_file", side_effect=AssertionError("Cache miss")):
            cached = processor.process([SimpleNamespace(name=str(path))])
        self.assertEqual([c.page_content for c in cached], [c.page_content for c in chunks])
        self.assertEqual(app._get_file_hashes([str(path)]), app._get_file_hashes([path]))

    def test_docling_converts_text_and_word(self):
        from docx import Document
        from document_processor.file_handler import DocumentProcessor

        text = self.root / "report.txt"
        text.write_text("The report records 42 cases.", encoding="utf-8")
        word = self.root / "report.docx"
        document = Document()
        document.add_heading("Cases", 1)
        document.add_paragraph("The Word report records 53 cases.")
        document.save(word)
        processor = DocumentProcessor()
        for path, expected in [(text, "42 cases"), (word, "53 cases")]:
            with self.subTest(format=path.suffix):
                chunks = processor.process([path])
                self.assertTrue(chunks)
                self.assertIn(expected, " ".join(chunk.page_content for chunk in chunks))

    def test_torch_and_torchvision_native_operators(self):
        import torch
        from torchvision.ops import nms

        boxes = torch.tensor([[0.0, 0.0, 10.0, 10.0], [0.0, 0.0, 9.0, 9.0]])
        scores = torch.tensor([0.9, 0.8])
        self.assertEqual(nms(boxes, scores, 0.5).tolist(), [0])

    def test_hybrid_retrieval_with_real_chroma_and_bm25(self):
        from langchain_core.documents import Document
        from langchain_core.embeddings import DeterministicFakeEmbedding
        from retriever.builder import RetrieverBuilder

        embeddings = DeterministicFakeEmbedding(size=16)
        with patch("retriever.builder.get_embedding", return_value=embeddings):
            builder = RetrieverBuilder()
        documents = [
            Document(page_content="The incident report records 42 cases."),
            Document(page_content="The maintenance plan schedules monthly inspections."),
        ]
        retriever = builder.build_hybrid_retriever(documents)
        result = retriever.invoke("incident cases")
        self.assertTrue(result)
        self.assertTrue(any("42 cases" in doc.page_content for doc in result))
        self.assertTrue((self.root / "chroma" / "chroma.sqlite3").exists())
        # Release Chroma's file handles before removing a temporary Windows directory.
        from chromadb.api.client import SharedSystemClient
        for system in list(SharedSystemClient._identifier_to_system.values()):
            system.stop()
        SharedSystemClient.clear_system_cache()

    def test_langgraph_flow_uses_real_groq_message_contract(self):
        from agents.workflow import AgentWorkflow
        from langchain_core.documents import Document
        from langchain_core.messages import AIMessage
        from langchain_groq import ChatGroq

        def respond(prompt, *args, **kwargs):
            if "Classify how well" in prompt:
                return AIMessage(content="CAN_ANSWER")
            if "verify the accuracy" in prompt:
                return AIMessage(content=(
                    "Supported: YES\nUnsupported Claims: []\nContradictions: []\n"
                    "Relevant: YES\nAdditional Details: Evidence found"
                ))
            return AIMessage(content="The report records 42 cases.")

        documents = [Document(page_content="The report records 42 cases.")]
        retriever = SimpleNamespace(invoke=lambda question: documents)
        with patch.object(ChatGroq, "invoke", side_effect=respond) as inference:
            result = AgentWorkflow().full_pipeline("How many cases?", retriever)
        self.assertEqual(result["draft_answer"], "The report records 42 cases.")
        self.assertIn("**Supported:** YES", result["verification_report"])
        self.assertEqual(inference.call_count, 3)

    def test_gradio_interface_builds_with_new_launch_parameters(self):
        import app
        import gradio as gr

        with patch.object(gr.Blocks, "launch", autospec=True) as launch:
            app.main()
        self.assertIn("theme", launch.call_args.kwargs)
        self.assertIn("css", launch.call_args.kwargs)
        self.assertIn("js", launch.call_args.kwargs)
        demo = launch.call_args.args[0]
        load_example = next(f.fn for f in demo.fns.values() if f.fn.__name__ == "load_example")
        paths, question = load_example("Google 2024 Environmental Report")
        self.assertTrue(question)
        self.assertTrue(paths)
        self.assertTrue(all(Path(path).is_file() for path in paths))
        self.assertEqual(launch.call_args.kwargs["server_name"], "127.0.0.1")
        self.assertEqual(launch.call_args.kwargs["server_port"], 5000)
        self.assertFalse(launch.call_args.kwargs["share"])

    def test_groq_404_aborts_workflow_without_marking_document_irrelevant(self):
        import httpx
        from groq import NotFoundError
        from agents.workflow import AgentWorkflow
        from config.llm import GroqServiceError
        from langchain_core.documents import Document
        from langchain_groq import ChatGroq

        request = httpx.Request("POST", "https://api.groq.com/openai/v1/chat/completions")
        error = NotFoundError(
            "Model unavailable",
            response=httpx.Response(404, request=request),
            body={"error": {"code": "model_not_found"}},
        )
        retriever = SimpleNamespace(invoke=lambda question: [Document(page_content="A synthetic report.")])
        workflow = AgentWorkflow()
        with patch.object(ChatGroq, "invoke", side_effect=error), patch.object(workflow.researcher, "generate") as research:
            with self.assertRaisesRegex(GroqServiceError, "GROQ_MODEL"):
                workflow.full_pipeline("What is this report about?", retriever)
            research.assert_not_called()

    def test_local_embedding_reuses_encoder_and_retrieves_spanish(self):
        import math
        from retriever.embeddings import get_embedding

        embedding = get_embedding()
        self.assertIs(embedding, get_embedding())
        vectors = embedding.embed_documents([
            "Para restablecer la contraseña de acceso, solicite un enlace de recuperación al correo registrado.",
            "La limpieza y mantenimiento de las cámaras se programan mensualmente.",
        ])
        query = embedding.embed_query("¿Cómo recuperar la contraseña de acceso?")
        self.assertEqual(len(query), 384)
        for vector in [query, *vectors]:
            self.assertAlmostEqual(math.sqrt(sum(x*x for x in vector)), 1.0, places=4)
        scores = [sum(a*b for a, b in zip(query, vector)) for vector in vectors]
        self.assertGreater(scores[0], scores[1])

    def test_long_document_preserves_tail_and_respects_token_limit(self):
        from document_processor.file_handler import DocumentProcessor
        from retriever.embeddings import get_tokenizer

        path = self.root / "large.md"
        body = " ".join(f"El procedimiento número {i} registra incidentes de seguridad." for i in range(500))
        path.write_text(f"# Procedimientos\n\nMARCADORINICIO {body} MARCADORFINAL", encoding="utf-8")
        chunks = DocumentProcessor().process([path])
        self.assertGreater(len(chunks), 10)
        self.assertIn("MARCADORINICIO", chunks[0].page_content)
        self.assertIn("MARCADORFINAL", chunks[-1].page_content)
        tokenizer = get_tokenizer()
        for chunk in chunks:
            self.assertLessEqual(len(tokenizer.tokenize(chunk.page_content)), self.settings.CHUNK_SIZE_TOKENS)
            self.assertLessEqual(len(tokenizer.encode("passage: " + chunk.page_content)), 512)
            self.assertEqual(chunk.metadata["source"], "large.md")
            self.assertEqual(chunk.metadata["Header 1"], "Procedimientos")

    def test_model_profiles_apply_retrieval_prompts(self):
        from retriever.embeddings import _load_embedding

        for name in ["intfloat/multilingual-e5-small", "BAAI/bge-m3"]:
            with self.subTest(model=name), patch("retriever.embeddings.HuggingFaceEmbeddings") as factory:
                _load_embedding.__wrapped__(name, "cpu", str(self.root), 16)
                kwargs = factory.call_args.kwargs
                self.assertTrue(kwargs["encode_kwargs"]["normalize_embeddings"])
                self.assertTrue(kwargs["query_encode_kwargs"]["normalize_embeddings"])
                if name == "intfloat/multilingual-e5-small":
                    self.assertEqual(kwargs["encode_kwargs"]["prompt"], "passage: ")
                    self.assertEqual(kwargs["query_encode_kwargs"]["prompt"], "query: ")
                else:
                    self.assertNotIn("prompt", kwargs["encode_kwargs"])
                    self.assertNotIn("prompt", kwargs["query_encode_kwargs"])

    def test_model_and_chunk_changes_isolate_index_and_cache(self):
        from config.settings import Settings
        from pydantic import ValidationError

        values = {"GROQ_API_KEY": "test-groq-key", "_env_file": None}
        e5 = Settings(**values)
        bge = Settings(**values, EMBEDDING_MODEL="BAAI/bge-m3", CHUNK_SIZE_TOKENS=1024)
        smaller = Settings(**values, CHUNK_SIZE_TOKENS=256)
        self.assertNotEqual(e5.embedding_collection_name, bge.embedding_collection_name)
        self.assertNotEqual(e5.embedding_fingerprint, smaller.embedding_fingerprint)
        with self.assertRaises(ValidationError):
            Settings(**values, CHUNK_SIZE_TOKENS=512)
        with self.assertRaises(ValidationError):
            Settings(**values, CHUNK_SIZE_TOKENS=64, CHUNK_OVERLAP_TOKENS=64)


if __name__ == "__main__":
    unittest.main()
